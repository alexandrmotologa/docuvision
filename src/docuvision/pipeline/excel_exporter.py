"""Professional multi-sheet Excel (.xlsx) workbook generator for extraction artifacts."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List, Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from docuvision.core.models import DocumentResult, Table


class ExcelExporter:
    """Exports DocumentResult objects into structured, styled multi-sheet Excel workbooks."""

    def __init__(self) -> None:
        # Styles definition
        self.font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.font_title = Font(name="Calibri", size=14, bold=True, color="1E293B")
        self.font_section = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
        self.font_bold = Font(name="Calibri", size=10, bold=True)
        self.font_normal = Font(name="Calibri", size=10)

        self.fill_primary = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        self.fill_accent = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        self.fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        self.fill_passed = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        self.fill_failed = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        self.border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

    def _auto_adjust_columns(self, ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        """Dynamically fit column widths to text length."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and len(val_str) < 60:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    def export_workbook(self, result: DocumentResult) -> openpyxl.Workbook:
        """Create openpyxl Workbook with Overview, Tables, and Audit sheets."""
        wb = openpyxl.Workbook()

        # ---------------------------------------------------------
        # Sheet 1: Document Overview
        # ---------------------------------------------------------
        ws_overview = wb.active
        ws_overview.title = "Document Overview"
        ws_overview.views.sheetView[0].showGridLines = True

        ws_overview["A1"] = "DocuVision Intelligence Summary"
        ws_overview["A1"].font = self.font_title
        ws_overview["A2"] = f"Extraction Schema: {result.schema_name.upper()}"
        ws_overview["A2"].font = self.font_section

        # Metadata Section
        ws_overview["A4"] = "Telemetry & Metadata"
        ws_overview["A4"].font = self.font_bold
        meta_rows = [
            ("File Name", result.metadata.filename),
            ("Total Pages", result.metadata.total_pages),
            ("File Size (KB)", round(result.metadata.file_size_bytes / 1024, 2)),
            ("SHA-256 Checksum", result.metadata.sha256_hash),
            ("Processing Latency (ms)", result.latency_ms),
            ("Validation Status", "PASSED" if result.validation_passed else "WARNINGS DETECTED"),
        ]

        for idx, (prop, val) in enumerate(meta_rows, start=5):
            cell_prop = ws_overview.cell(row=idx, column=1, value=prop)
            cell_val = ws_overview.cell(row=idx, column=2, value=val)
            cell_prop.font = self.font_bold
            cell_val.font = self.font_normal
            cell_prop.border = self.border_thin
            cell_val.border = self.border_thin

        # Key Fields Section
        start_row = len(meta_rows) + 7
        ws_overview.cell(row=start_row, column=1, value="Key Structured Fields").font = self.font_section
        field_start = start_row + 1

        scalar_fields = [(k, v) for k, v in result.structured_data.items() if not isinstance(v, (list, dict))]
        if scalar_fields:
            for idx, (k, v) in enumerate(scalar_fields, start=field_start):
                cp = ws_overview.cell(row=idx, column=1, value=k)
                cv = ws_overview.cell(row=idx, column=2, value=str(v) if v is not None else "")
                cp.font = self.font_bold
                cv.font = self.font_normal
                cp.border = self.border_thin
                cv.border = self.border_thin

        self._auto_adjust_columns(ws_overview)

        # ---------------------------------------------------------
        # Sheet 2: Extracted Tables
        # ---------------------------------------------------------
        ws_tables = wb.create_sheet(title="Extracted Tables")
        ws_tables.views.sheetView[0].showGridLines = True

        current_row = 1
        if result.tables:
            for t_idx, tab in enumerate(result.tables):
                ws_tables.cell(row=current_row, column=1, value=f"Table {t_idx + 1} (Page {tab.page_number} - {tab.rows} Rows x {tab.cols} Cols)").font = self.font_section
                current_row += 1

                for r_idx, row in enumerate(tab.matrix):
                    is_header = r_idx == 0
                    for c_idx, val in enumerate(row):
                        cell = ws_tables.cell(row=current_row, column=c_idx + 1, value=val)
                        cell.border = self.border_thin

                        if is_header:
                            cell.font = self.font_header
                            cell.fill = self.fill_accent
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        else:
                            cell.font = self.font_normal
                            if r_idx % 2 == 1:
                                cell.fill = self.fill_zebra
                    current_row += 1
                current_row += 2  # Spacer between tables
        else:
            ws_tables.cell(row=1, column=1, value="No tabular regions detected.").font = self.font_normal

        self._auto_adjust_columns(ws_tables)

        # ---------------------------------------------------------
        # Sheet 3: Audit & Validation
        # ---------------------------------------------------------
        ws_audit = wb.create_sheet(title="Audit & Validation")
        ws_audit.views.sheetView[0].showGridLines = True

        ws_audit["A1"] = "Mathematical Reconciliation & Compliance Report"
        ws_audit["A1"].font = self.font_title

        status_cell = ws_audit["A3"]
        status_cell.value = "PASSED: All line items and totals validated consistently" if result.validation_passed else "WARNINGS: Discrepancies detected"
        status_cell.font = self.font_bold
        status_cell.fill = self.fill_passed if result.validation_passed else self.fill_failed
        status_cell.alignment = Alignment(horizontal="left", vertical="center")

        ws_audit["A5"] = "Discrepancy Details:"
        ws_audit["A5"].font = self.font_bold

        if result.validation_errors:
            for idx, err in enumerate(result.validation_errors, start=6):
                c = ws_audit.cell(row=idx, column=1, value=f"• {err}")
                c.font = self.font_normal
        else:
            ws_audit["A6"] = "No mathematical discrepancies or schema validation errors found."
            ws_audit["A6"].font = self.font_normal

        self._auto_adjust_columns(ws_audit)

        return wb

    def export_to_bytes(self, result: DocumentResult) -> bytes:
        """Generate workbook and return raw .xlsx bytes."""
        wb = self.export_workbook(result)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def export_to_file(self, result: DocumentResult, destination: Union[str, Path]) -> Path:
        """Write multi-sheet Excel file to disk."""
        dest_path = Path(destination)
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        wb = self.export_workbook(result)
        wb.save(str(dest_path))
        return dest_path
